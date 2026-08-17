"""Parse the macro-free authoring workbook."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from push_challenges_core.errors import Diagnostic

from .normalize import AuthoringDocument, AuthoringRow
from .workbook import (
    CHALLENGE_HEADERS,
    FLAG_HEADERS,
    HINT_HEADERS,
    SOLUTION_HEADERS,
)

REQUIRED_SHEETS = {
    "Challenges": CHALLENGE_HEADERS,
    "Flags": FLAG_HEADERS,
    "Solutions": SOLUTION_HEADERS,
    "Hints": HINT_HEADERS,
}
FLAG_TYPES = {"Static": "static", "Regex": "regex"}
FLAG_MATCHING = {
    "Case Sensitive": "case_sensitive",
    "Case Insensitive": "case_insensitive",
}


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _sheet_rows(
    sheet,
    headers: tuple[str, ...],
    *,
    default_columns: set[str],
) -> tuple[list[tuple[int, dict[str, object]]], Diagnostic | None]:
    actual = tuple(sheet.cell(1, index).value for index in range(1, len(headers) + 1))
    if actual != headers:
        return [], Diagnostic(
            "invalid_xlsx_headers",
            "Expected exact headers: " + ", ".join(headers),
            f"{sheet.title}!1",
            True,
        )
    rows: list[tuple[int, dict[str, object]]] = []
    meaningful = tuple(
        index for index, header in enumerate(headers) if header not in default_columns
    )
    for row_number in range(2, sheet.max_row + 1):
        values = tuple(
            sheet.cell(row_number, column).value
            for column in range(1, len(headers) + 1)
        )
        if all(_blank(values[index]) for index in meaningful):
            continue
        rows.append((row_number, dict(zip(headers, values, strict=True))))
    return rows, None


def _diagnostic(
    code: str, message: str, sheet: str, row: int, *, blocking: bool = True
) -> Diagnostic:
    return Diagnostic(code, message, f"{sheet}!A{row}", blocking)


def load_xlsx(path: Path) -> AuthoringDocument:
    diagnostics: list[Diagnostic] = []
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except (OSError, ValueError) as error:
        return AuthoringDocument(
            (),
            (Diagnostic("xlsx_read_failed", str(error), str(path), True),),
        )

    missing = tuple(name for name in REQUIRED_SHEETS if name not in workbook.sheetnames)
    if missing:
        return AuthoringDocument(
            (),
            (
                Diagnostic(
                    "missing_xlsx_sheet",
                    "Missing required sheet(s): " + ", ".join(missing),
                    str(path),
                    True,
                ),
            ),
        )

    parsed: dict[str, list[tuple[int, dict[str, object]]]] = {}
    defaults = {
        "Challenges": {"State", "Type", "Logic", "Max Attempts"},
        "Flags": set(),
        "Solutions": {"State"},
        "Hints": {"Cost"},
    }
    for name, headers in REQUIRED_SHEETS.items():
        rows, error = _sheet_rows(
            workbook[name], headers, default_columns=defaults[name]
        )
        if error is not None:
            diagnostics.append(error)
        parsed[name] = rows
    if any(item.blocking for item in diagnostics):
        return AuthoringDocument((), tuple(diagnostics))

    challenge_rows: list[AuthoringRow] = []
    values_by_name: dict[str, dict[str, object]] = {}
    source_by_name: dict[str, str] = {}
    for row_number, values in parsed["Challenges"]:
        name = "" if values["Name"] is None else str(values["Name"])
        values["Flags"] = []
        values["Hints"] = []
        values["Solution"] = ""
        values["Solution State"] = ""
        source = f"Challenges row {row_number}"
        challenge_rows.append(AuthoringRow(source, values))
        values_by_name.setdefault(name, values)
        source_by_name.setdefault(name, source)

    for row_number, values in parsed["Flags"]:
        challenge = "" if values["Challenge"] is None else str(values["Challenge"])
        if challenge not in values_by_name:
            diagnostics.append(
                _diagnostic(
                    "orphan_flag",
                    f"Flag references missing input challenge {challenge!r}",
                    "Flags",
                    row_number,
                )
            )
            continue
        values_by_name[challenge]["Flags"].append(
            {
                "type": FLAG_TYPES.get(values["Flag Type"], values["Flag Type"]),
                "content": values["Content"],
                "data": FLAG_MATCHING.get(values["Matching"], values["Matching"]),
            }
        )

    solution_rows: defaultdict[str, list[tuple[int, dict[str, object]]]] = defaultdict(
        list
    )
    for row_number, values in parsed["Solutions"]:
        challenge = "" if values["Challenge"] is None else str(values["Challenge"])
        if challenge not in values_by_name:
            diagnostics.append(
                _diagnostic(
                    "orphan_solution",
                    f"Solution references missing input challenge {challenge!r}",
                    "Solutions",
                    row_number,
                )
            )
            continue
        solution_rows[challenge].append((row_number, values))
    for challenge, rows in solution_rows.items():
        if len(rows) > 1:
            diagnostics.append(
                _diagnostic(
                    "duplicate_solution_row",
                    f"Multiple solution rows reference {challenge!r}",
                    "Solutions",
                    rows[1][0],
                )
            )
            continue
        _, values = rows[0]
        values_by_name[challenge]["Solution"] = values["Content"]
        values_by_name[challenge]["Solution State"] = values["State"]

    for row_number, values in parsed["Hints"]:
        challenge = "" if values["Challenge"] is None else str(values["Challenge"])
        if challenge not in values_by_name:
            diagnostics.append(
                _diagnostic(
                    "orphan_hint_ignored",
                    f"Hint references missing input challenge {challenge!r}; ignored",
                    "Hints",
                    row_number,
                    blocking=False,
                )
            )
            continue
        values_by_name[challenge]["Hints"].append(
            {
                "title": values["Title"],
                "content": values["Hint"],
                "cost": values["Cost"],
                "required_hints": values["Required Hints"] or "",
            }
        )

    return AuthoringDocument(tuple(challenge_rows), tuple(diagnostics))
