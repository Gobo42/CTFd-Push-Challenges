"""Create the macro-free challenge authoring workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

CHALLENGE_HEADERS = (
    "Name",
    "Description",
    "Value",
    "Category",
    "State",
    "Type",
    "Tags",
    "Next",
    "Logic",
    "Max Attempts",
    "Attribution",
    "Connection Info",
)
FLAG_HEADERS = ("Challenge", "Flag Type", "Content", "Matching")
SOLUTION_HEADERS = ("Challenge", "Content", "State")
HINT_HEADERS = ("Challenge", "Title", "Hint", "Cost", "Required Hints")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")


def _style_sheet(sheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _table(sheet, name: str, last_row: int, last_column: int) -> None:
    reference = f"A1:{get_column_letter(last_column)}{last_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _list_validation(sheet, formula: str, cells: str) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Choose a value from the dropdown."
    validation.errorTitle = "Invalid value"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(cells)


def _whole_number_validation(
    sheet, cells: str, *, minimum: int, allow_blank: bool
) -> None:
    validation = DataValidation(
        type="whole",
        operator="greaterThanOrEqual",
        formula1=str(minimum),
        allow_blank=allow_blank,
    )
    validation.error = f"Enter a whole number greater than or equal to {minimum}."
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(cells)


def _add_defined_name(workbook: Workbook, name: str, formula: str) -> None:
    workbook.defined_names.add(DefinedName(name, attr_text=formula))


def _create_lists(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Lists")
    columns = {
        "A": ("Challenge States", "Hidden", "Visible", "Locked"),
        "B": ("Challenge Types", "Standard", "Dynamic"),
        "C": ("Flag Logic", "Require Any Flag", "Require All Flags"),
        "D": ("Flag Types", "Static", "Regex"),
        "E": ("Flag Matching", "Case Sensitive", "Case Insensitive"),
        "F": ("Solution States", "Hidden", "Visible", "Solved"),
    }
    names = {
        "A": "ChallengeStates",
        "B": "ChallengeTypes",
        "C": "FlagLogic",
        "D": "FlagTypes",
        "E": "FlagMatching",
        "F": "SolutionStates",
    }
    for column, values in columns.items():
        for row, value in enumerate(values, start=1):
            sheet[f"{column}{row}"] = value
        _add_defined_name(
            workbook,
            names[column],
            f"'Lists'!${column}$2:${column}${len(values)}",
        )
    _add_defined_name(
        workbook,
        "ChallengeNames",
        "OFFSET('Challenges'!$A$2,0,0,"
        "MAX(1,COUNTA('Challenges'!$A:$A)-1),1)",
    )
    sheet.sheet_state = "hidden"


def _create_instructions(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "Instructions"
    lines = (
        ("Push Challenges workbook", "Author challenges here, then validate/export."),
        (
            "Workflow",
            "Complete Challenges and related sheets. Run validate_challenges.py "
            "to review decoded content and create a .push.json manifest.",
        ),
        (
            "Tags",
            "Start every tag with #. Spaces belong to the current tag until the "
            "next #, for example: #beginner #web authentication.",
        ),
        (
            "Solutions",
            "No row or blank Content preserves the existing solution. Exact "
            "case-sensitive REMOVE clears solution text and hides it without "
            "deleting its record or files.",
        ),
        (
            "Hints",
            "Required Hints uses comma-separated 1-based numbers of earlier hints "
            "for that challenge. Invalid requirements are ignored with a warning.",
        ),
        (
            "Limitations",
            "Standard challenges are supported. Dynamic is reserved for future use "
            "and currently fails validation. Files remain manual.",
        ),
    )
    for row, values in enumerate(lines, start=1):
        sheet.cell(row, 1, values[0])
        sheet.cell(row, 2, values[1])
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def create_workbook(path: Path, author_rows: int = 500) -> None:
    if author_rows < 1:
        raise ValueError("author_rows must be at least 1")
    workbook = Workbook()
    _create_instructions(workbook)

    challenges = workbook.create_sheet("Challenges")
    challenges.append(CHALLENGE_HEADERS)
    for row in range(2, author_rows + 2):
        challenges.cell(row, 5, "Hidden")
        challenges.cell(row, 6, "Standard")
        challenges.cell(row, 9, "Require Any Flag")
        challenges.cell(row, 10, 0)
    _style_sheet(
        challenges,
        (28, 70, 12, 20, 14, 14, 38, 28, 22, 16, 28, 42),
    )
    _list_validation(challenges, "=ChallengeStates", f"E2:E{author_rows + 1}")
    _list_validation(challenges, "=ChallengeTypes", f"F2:F{author_rows + 1}")
    _list_validation(challenges, "=FlagLogic", f"I2:I{author_rows + 1}")
    _whole_number_validation(
        challenges, f"J2:J{author_rows + 1}", minimum=0, allow_blank=False
    )
    challenges.conditional_formatting.add(
        f"F2:F{author_rows + 1}",
        FormulaRule(formula=['F2="Dynamic"'], fill=WARNING_FILL),
    )
    _table(challenges, "ChallengesTable", author_rows + 1, len(CHALLENGE_HEADERS))

    flags = workbook.create_sheet("Flags")
    flags.append(FLAG_HEADERS)
    for _ in range(author_rows):
        flags.append((None,) * len(FLAG_HEADERS))
    _style_sheet(flags, (28, 16, 60, 20))
    _list_validation(flags, "=ChallengeNames", f"A2:A{author_rows + 1}")
    _list_validation(flags, "=FlagTypes", f"B2:B{author_rows + 1}")
    _list_validation(flags, "=FlagMatching", f"D2:D{author_rows + 1}")
    _table(flags, "FlagsTable", author_rows + 1, len(FLAG_HEADERS))

    solutions = workbook.create_sheet("Solutions")
    solutions.append(SOLUTION_HEADERS)
    for row in range(2, author_rows + 2):
        solutions.cell(row, 3, "Hidden")
    _style_sheet(solutions, (28, 80, 16))
    _list_validation(solutions, "=ChallengeNames", f"A2:A{author_rows + 1}")
    _list_validation(solutions, "=SolutionStates", f"C2:C{author_rows + 1}")
    _table(solutions, "SolutionsTable", author_rows + 1, len(SOLUTION_HEADERS))

    hints = workbook.create_sheet("Hints")
    hints.append(HINT_HEADERS)
    for row in range(2, author_rows + 2):
        hints.cell(row, 4, 0)
    _style_sheet(hints, (28, 30, 80, 12, 24))
    _list_validation(hints, "=ChallengeNames", f"A2:A{author_rows + 1}")
    _whole_number_validation(
        hints, f"D2:D{author_rows + 1}", minimum=0, allow_blank=False
    )
    _table(hints, "HintsTable", author_rows + 1, len(HINT_HEADERS))

    _create_lists(workbook)
    workbook.save(path)
