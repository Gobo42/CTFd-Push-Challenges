import importlib.util
import sys
from copy import deepcopy
from pathlib import Path

import pytest
from conftest import challenge_data, manifest_data
from CTFd.models import Awards, Challenges, Flags, Solves, Tags, Users, db
from CTFd.plugins.dynamic_challenges import DynamicChallenge
from sqlalchemy import event
from test_preflight import importer_module, persisted_challenge

from push_challenges_core.manifest import validate_manifest_data


class CustomChallenge(Challenges):
    __tablename__ = None
    __mapper_args__ = {"polymorphic_identity": "push_test_custom"}


def scoring_data(name="Dynamic", *, bonus=None, function="linear"):
    data = challenge_data(name)
    data.update(
        type="dynamic",
        scoring={"function": function, "minimum": 20, "decay": 10},
        bonus=bonus or {"mode": "preserve", "tiers": []},
    )
    return data


def scoring_manifest(*challenges):
    data = manifest_data(*deepcopy(challenges))
    data["version"] = 2
    for challenge in data["challenges"]:
        challenge.setdefault("scoring", None)
        challenge.setdefault("bonus", {"mode": "preserve", "tiers": []})
    manifest, diagnostics = validate_manifest_data(data)
    assert manifest is not None, diagnostics
    return manifest


@pytest.mark.parametrize("function", ["linear", "logarithmic"])
def test_native_dynamic_create_and_update(app, function):
    importer = importer_module()
    spec = scoring_data(function=function)
    importer.apply_plan(importer.build_plan(scoring_manifest(spec)))
    challenge = DynamicChallenge.query.one()
    assert (challenge.type, challenge.function, challenge.initial) == (
        "dynamic",
        function,
        100,
    )
    assert (challenge.minimum, challenge.decay, challenge.value) == (20, 10, 100)
    challenge_id = challenge.id
    spec["value"] = 200
    spec["scoring"]["minimum"] = 50
    importer.apply_plan(importer.build_plan(scoring_manifest(spec)))
    assert DynamicChallenge.query.one().id == challenge_id
    assert challenge.initial == challenge.value == 200


def test_mismatch_skips_all_managed_data_and_keeps_next_target(app):
    importer = importer_module()
    existing = persisted_challenge("Existing")
    target = persisted_challenge("Target")
    existing.next_id = target.id
    db.session.add_all(
        [
            Tags(challenge_id=existing.id, value="old"),
        ]
    )
    db.session.execute(
        Flags.__table__.insert(),
        [{"challenge_id": existing.id, "type": "static", "content": "old", "data": ""}],
    )
    db.session.commit()
    # Even bonus requests on skipped rows do not require Monthly Game.
    skipped = scoring_data("Existing", bonus={"mode": "disabled", "tiers": []})
    new = challenge_data("New", next_name="Existing")
    plan = importer.build_plan(scoring_manifest(skipped, new))
    assert len(plan.skipped) == 1
    assert plan.projected_next["Existing"] == "Target"
    importer.apply_plan(plan)
    assert existing.description == "Existing"
    assert existing.value == 50
    assert existing.next_id == target.id
    assert Tags.query.filter_by(challenge_id=existing.id).one().value == "old"
    assert Flags.query.filter_by(challenge_id=existing.id).one().content == "old"
    assert Challenges.query.filter_by(name="New").one().next_id == existing.id


def test_cycle_through_skipped_row_is_rejected(app):
    importer = importer_module()
    existing = persisted_challenge("Existing")
    target = persisted_challenge("Target", next_id=existing.id)
    with pytest.raises(importer.ImportPreflightError, match="Circular Next"):
        importer.build_plan(
            scoring_manifest(
                scoring_data("Target"), challenge_data("Existing", next_name="Target")
            )
        )
    assert target.next_id == existing.id


def test_bonus_dependency_checked_before_writes(app):
    importer = importer_module()
    app.extensions.pop("monthly_game_bonuses", None)
    spec = scoring_data(
        bonus={"mode": "absolute", "tiers": [{"start": 1, "end": 3, "amount": "25"}]}
    )
    with pytest.raises(importer.ImportPreflightError, match="Monthly Game"):
        importer.build_plan(scoring_manifest(spec))
    assert Challenges.query.count() == 0


def test_outdated_bonus_service_blocks_preflight(app):
    importer = importer_module()
    app.extensions["monthly_game_bonuses"] = object()
    with pytest.raises(importer.ImportPreflightError, match="install or update"):
        importer.build_plan(
            scoring_manifest(scoring_data(bonus={"mode": "disabled", "tiers": []}))
        )
    assert Challenges.query.count() == 0


def test_dynamic_to_standard_is_also_skipped(app):
    importer = importer_module()
    challenge = DynamicChallenge(
        name="Existing",
        type="dynamic",
        initial=100,
        minimum=20,
        decay=10,
        function="logarithmic",
    )
    db.session.add(challenge)
    db.session.commit()
    plan = importer.build_plan(scoring_manifest(challenge_data("Existing")))
    assert len(plan.skipped) == 1
    assert plan.challenges == ()
    importer.apply_plan(plan)
    assert challenge.initial == challenge.value == 100
    assert challenge.function == "logarithmic"


def test_unsupported_custom_type_still_aborts(app):
    importer = importer_module()
    challenge = CustomChallenge(name="Custom", type="push_test_custom", value=10)
    db.session.add(challenge)
    db.session.commit()
    with pytest.raises(
        importer.ImportPreflightError, match="unsupported challenge type"
    ):
        importer.build_plan(scoring_manifest(challenge_data("Custom")))


@pytest.mark.parametrize("apply", [False, True])
def test_cli_mismatch_has_final_manual_callout(app, runner, tmp_path, apply):
    persisted_challenge("Existing")
    manifest = scoring_manifest(scoring_data("Existing"), challenge_data("New"))
    from push_challenges_core.manifest import dump_manifest

    path = tmp_path / "challenges.json"
    dump_manifest(manifest, path)
    args = ["push-challenges", "import", str(path)] + (["--apply"] if apply else [])
    result = runner.invoke(args=args)
    assert result.exit_code == 0, result.output
    assert "skipped: 1" in result.output
    callout = result.output.index("MANUAL ACTION REQUIRED")
    assert callout > result.output.index("APPLIED:" if apply else "DRY RUN:")
    assert "Existing" in result.output[callout:]
    assert "standard -> dynamic" in result.output[callout:]
    assert "manually" in result.output[callout:]
    assert Challenges.query.filter_by(name="New").count() == int(apply)


def test_dynamic_failure_rolls_back_joined_table_and_standard(app):
    importer = importer_module()
    plan = importer.build_plan(scoring_manifest(scoring_data(), challenge_data("New")))

    def fail(stage):
        if stage == "after_next":
            raise RuntimeError("injected failure")

    with pytest.raises(importer.ImportApplyError, match="rolled back"):
        importer.apply_plan(plan, stage_hook=fail)
    assert Challenges.query.count() == DynamicChallenge.query.count() == 0


@pytest.mark.parametrize(
    ("function", "expected"), [("linear", 80), ("logarithmic", 97)]
)
def test_dynamic_overwrite_recalculates_native_value_from_existing_solves(
    app, function, expected, monkeypatch
):
    from CTFd.plugins.dynamic_challenges import decay

    importer = importer_module()
    challenge = DynamicChallenge(
        name="Dynamic",
        type="dynamic",
        initial=100,
        minimum=20,
        decay=10,
        function=function,
    )
    db.session.add(challenge)
    db.session.flush()
    for index in range(3):
        user = Users(name=f"user{index}", email=f"user{index}@example.com")
        db.session.add(user)
        db.session.flush()
        db.session.add(
            Solves(
                user_id=user.id,
                challenge_id=challenge.id,
                ip="127.0.0.1",
                provided="flag",
            )
        )
    db.session.commit()
    # Test applications do not set up a competition mode.
    monkeypatch.setattr(decay, "get_model", lambda: Users)
    importer.apply_plan(
        importer.build_plan(scoring_manifest(scoring_data(function=function)))
    )
    assert challenge.value == expected
    assert Solves.query.count() == 3


def test_apply_refuses_type_race_without_partial_writes(app):
    importer = importer_module()
    challenge = persisted_challenge("Existing")
    plan = importer.build_plan(
        scoring_manifest(challenge_data("Existing"), scoring_data())
    )
    # Simulate another transaction changing the discriminator.
    db.session.execute(
        Challenges.__table__.update()
        .where(Challenges.id == challenge.id)
        .values(type="dynamic")
    )
    db.session.commit()
    with pytest.raises(importer.ImportApplyError, match="changed type"):
        importer.apply_plan(plan)
    assert DynamicChallenge.query.count() == 0


@pytest.fixture
def monthly(app):
    name = "CTFd.plugins.monthly-game"
    root = Path(__file__).resolve().parents[3] / "monthly_game"
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(
            name, root / "__init__.py", submodule_search_locations=[str(root)]
        )
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    module = sys.modules[name]
    app._got_first_request = False
    module.load(app)
    try:
        yield sys.modules[name + ".models"]
    finally:
        sys.modules[name + ".ledger"].unregister_score_hooks()
        sys.modules[name + ".page_variables"].unregister_page_variable_session_hooks()


def test_bonus_import_uses_shared_service_one_commit_and_preserves_awards(app, monthly):
    importer = importer_module()
    spec = challenge_data("Bonus")
    spec["bonus"] = {
        "mode": "absolute",
        "tiers": [{"start": 1, "end": 5, "amount": "25"}],
    }
    plan = importer.build_plan(scoring_manifest(spec))
    commits = []
    session = db.session()
    event.listen(session, "after_commit", lambda _: commits.append(1))
    importer.apply_plan(plan)
    assert commits == [1]
    program = monthly.MonthlyGameBonusProgram.query.one()
    assert program.next_position == 1
    assert program.enabled
    user = Users(name="Player", email="player@example.com")
    db.session.add(user)
    db.session.flush()
    db.session.add(
        Solves(
            user_id=user.id,
            challenge_id=program.challenge_id,
            ip="127.0.0.1",
            provided="flag",
        )
    )
    db.session.commit()
    assert Awards.query.one().value == 25
    assert program.next_position == 2
    grant_id = monthly.MonthlyGameBonusGrant.query.one().id
    spec["bonus"] = {
        "mode": "percentage",
        "tiers": [{"start": 1, "end": 10, "amount": "12.999"}],
    }
    importer.apply_plan(importer.build_plan(scoring_manifest(spec)))
    assert program.next_position == 2
    assert str(monthly.MonthlyGameBonusTier.query.one().amount) == "12.99"
    assert Awards.query.one().value == 25
    assert monthly.MonthlyGameBonusGrant.query.one().id == grant_id
    spec["bonus"] = {"mode": "disabled", "tiers": []}
    importer.apply_plan(importer.build_plan(scoring_manifest(spec)))
    assert not program.enabled
    assert program.next_position == 2
    assert monthly.MonthlyGameBonusTier.query.count() == 1
    spec["bonus"] = {"mode": "preserve", "tiers": []}
    importer.apply_plan(importer.build_plan(scoring_manifest(spec)))
    assert not program.enabled


def test_bonus_failure_rolls_back_challenges_and_tiers(app, monthly):
    importer = importer_module()
    spec = scoring_data(
        bonus={"mode": "absolute", "tiers": [{"start": 1, "end": 3, "amount": "25"}]}
    )
    plan = importer.build_plan(scoring_manifest(spec, challenge_data("New")))

    def fail(stage):
        if stage == "after_bonus":
            raise RuntimeError("injected bonus failure")

    with pytest.raises(importer.ImportApplyError, match="rolled back"):
        importer.apply_plan(plan, stage_hook=fail)
    assert Challenges.query.count() == 0
    assert monthly.MonthlyGameBonusProgram.query.count() == 0
    assert monthly.MonthlyGameBonusTier.query.count() == 0


def test_scoring_workbook_to_manifest_to_bonus_import(
    app, monthly, runner, tmp_path, capsys
):
    from openpyxl import load_workbook

    from authoring.workbook import create_workbook
    from tools.validate_challenges import main as validate_main

    path = tmp_path / "scoring.xlsx"
    manifest_path = tmp_path / "scoring.json"
    create_workbook(path, author_rows=2)
    book = load_workbook(path)
    values = {
        "Name": "Workbook challenge",
        "Value": 100,
        "Category": "Web",
        "Type": "Dynamic",
        "Scoring Function": "Logarithmic",
        "Minimum Value": 20,
        "Decay": 10,
        "Bonus Mode": "Percentage",
        "Position": 25,
    }
    for header in book["Challenges"][1]:
        if header.value in values:
            book["Challenges"].cell(2, header.column, values[header.value])
    for column, value in enumerate(("Workbook challenge", 1, 3, 25.999), 1):
        book["Bonus Tiers"].cell(2, column, value)
    book.save(path)
    assert validate_main([str(path), "--output", str(manifest_path)]) == 0
    assert "25.99" in capsys.readouterr().out
    dry_run = runner.invoke(args=["push-challenges", "import", str(manifest_path)])
    assert dry_run.exit_code == 0, dry_run.output
    assert Challenges.query.count() == 0
    assert monthly.MonthlyGameBonusProgram.query.count() == 0
    result = runner.invoke(
        args=["push-challenges", "import", str(manifest_path), "--apply"]
    )
    assert result.exit_code == 0, result.output
    assert DynamicChallenge.query.one().function == "logarithmic"
    assert DynamicChallenge.query.one().position == 25
    assert monthly.MonthlyGameBonusProgram.query.one().calculation_mode == "percentage"
    assert str(monthly.MonthlyGameBonusTier.query.one().amount) == "25.99"
