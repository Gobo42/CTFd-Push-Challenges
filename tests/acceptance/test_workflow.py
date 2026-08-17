from CTFd.models import Challenges, Hints, Solutions
from openpyxl import load_workbook

from authoring.workbook import create_workbook
from tools.validate_challenges import main as validate_main


def set_row(sheet, row_number, values):
    for column, value in enumerate(values, start=1):
        sheet.cell(row_number, column, value)


def test_workbook_export_dry_run_apply_and_exact_name_overwrite(
    app, runner, tmp_path, capsys
):
    workbook_path = tmp_path / "challenges.xlsx"
    manifest_path = tmp_path / "challenges.push.json"
    create_workbook(workbook_path, author_rows=10)
    workbook = load_workbook(workbook_path)
    set_row(
        workbook["Challenges"],
        2,
        [
            "First",
            "Initial prompt",
            100,
            "Intro",
            "Visible",
            "Standard",
            "#beginner #web authentication",
            "Second",
            "Require All Flags",
            2,
            "Author",
            "nc host 31337",
        ],
    )
    set_row(
        workbook["Challenges"],
        3,
        [
            "Second",
            "Follow-up",
            200,
            "Intro",
            "Hidden",
            "Standard",
            "#beginner",
            "",
            "Require Any Flag",
            0,
            "",
            "",
        ],
    )
    set_row(
        workbook["Flags"],
        2,
        ["First", "Static", "flag{one}", "Case Sensitive"],
    )
    set_row(
        workbook["Flags"],
        3,
        ["First", "Regex", r"flag\{two\}", "Case Insensitive"],
    )
    set_row(
        workbook["Solutions"],
        2,
        ["First", "Walkthrough", "Solved"],
    )
    set_row(
        workbook["Hints"],
        2,
        ["First", "Clue one", "Read it", 0, ""],
    )
    set_row(
        workbook["Hints"],
        3,
        ["First", "Clue two", "Decode it", 5, "1"],
    )
    set_row(
        workbook["Hints"],
        4,
        ["Deleted", "Orphan", "Ignore it", 0, ""],
    )
    workbook.save(workbook_path)

    export_code = validate_main(
        [str(workbook_path), "--output", str(manifest_path)]
    )
    export_output = capsys.readouterr().out

    assert export_code == 0
    assert "orphan_hint_ignored" in export_output
    assert manifest_path.exists()

    dry_run = runner.invoke(
        args=["push-challenges", "import", str(manifest_path)]
    )

    assert dry_run.exit_code == 0
    assert "DRY RUN" in dry_run.output
    assert Challenges.query.count() == 0

    apply = runner.invoke(
        args=[
            "push-challenges",
            "import",
            str(manifest_path),
            "--apply",
        ]
    )

    assert apply.exit_code == 0
    assert Challenges.query.count() == 2
    linked = Challenges.query.filter_by(name="First").one()
    target = Challenges.query.filter_by(name="Second").one()
    assert linked.next_id == target.id
    assert linked.solution.content == "Walkthrough"
    assert linked.solution.state == "solved"
    retained_hint = Hints.query.filter_by(
        challenge_id=linked.id, title="Clue one"
    ).one()
    original_hint_id = retained_hint.id

    workbook = load_workbook(workbook_path)
    workbook["Challenges"]["B2"] = "Overwritten prompt"
    workbook["Hints"]["C2"] = "Updated clue"
    workbook.save(workbook_path)
    assert (
        validate_main([str(workbook_path), "--output", str(manifest_path)])
        == 0
    )
    capsys.readouterr()
    second_apply = runner.invoke(
        args=[
            "push-challenges",
            "import",
            str(manifest_path),
            "--apply",
        ]
    )

    assert second_apply.exit_code == 0
    assert Challenges.query.count() == 2
    overwritten = Challenges.query.filter_by(name="First").one()
    assert overwritten.description == "Overwritten prompt"
    retained_hint = Hints.query.filter_by(
        challenge_id=overwritten.id, title="Clue one"
    ).one()
    assert retained_hint.id == original_hint_id
    assert retained_hint.content == "Updated clue"
    assert Solutions.query.count() == 1
