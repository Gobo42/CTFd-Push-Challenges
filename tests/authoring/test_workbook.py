import zipfile
from pathlib import Path

from openpyxl import load_workbook

from authoring.workbook import create_workbook

AUTHORING_SHEETS = ("Challenges", "Flags", "Solutions", "Hints")


def assert_table_filters_are_compatible(workbook):
    for sheet_name in AUTHORING_SHEETS:
        sheet = workbook[sheet_name]
        assert sheet.auto_filter.ref is None
        tables = tuple(sheet.tables.values())
        assert len(tables) == 1
        assert tables[0].autoFilter is not None
        assert tables[0].autoFilter.ref == tables[0].ref


def test_workbook_has_author_friendly_sheets_and_defaults(tmp_path):
    path = tmp_path / "challenges.xlsx"

    create_workbook(path, author_rows=20)
    workbook = load_workbook(path)

    assert workbook.sheetnames == [
        "Instructions",
        "Challenges",
        "Flags",
        "Solutions",
        "Hints",
        "Lists",
    ]
    assert workbook["Lists"].sheet_state == "hidden"
    assert workbook["Challenges"]["E2"].value == "Hidden"
    assert workbook["Challenges"]["F2"].value == "Standard"
    assert workbook["Challenges"]["I2"].value == "Require Any Flag"
    assert workbook["Challenges"]["J2"].value == 0
    assert workbook["Solutions"]["C2"].value == "Hidden"
    assert workbook["Hints"]["D2"].value == 0


def test_workbook_dropdowns_use_fixed_lists_and_challenge_names(tmp_path):
    path = tmp_path / "challenges.xlsx"
    create_workbook(path, author_rows=10)
    workbook = load_workbook(path)

    challenge_validations = tuple(
        workbook["Challenges"].data_validations.dataValidation
    )
    flag_validations = tuple(workbook["Flags"].data_validations.dataValidation)
    solution_validations = tuple(
        workbook["Solutions"].data_validations.dataValidation
    )
    hint_validations = tuple(workbook["Hints"].data_validations.dataValidation)

    assert any(item.formula1 == "=ChallengeStates" for item in challenge_validations)
    assert any(item.formula1 == "=ChallengeTypes" for item in challenge_validations)
    assert workbook["Lists"]["B3"].value == "Dynamic"
    assert any(item.formula1 == "=FlagLogic" for item in challenge_validations)
    assert any(item.formula1 == "=ChallengeNames" for item in flag_validations)
    assert any(item.formula1 == "=ChallengeNames" for item in solution_validations)
    assert any(item.formula1 == "=ChallengeNames" for item in hint_validations)


def test_workbook_is_macro_free(tmp_path):
    path = tmp_path / "challenges.xlsx"
    create_workbook(path, author_rows=5)

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()

    assert not any("vbaProject" in name for name in names)
    assert not any(name.endswith(".bin") for name in names)


def test_generated_workbook_uses_only_table_owned_filters(tmp_path):
    path = tmp_path / "challenges.xlsx"
    create_workbook(path, author_rows=5)

    assert_table_filters_are_compatible(load_workbook(path))


def test_distributed_template_uses_only_table_owned_filters():
    path = Path(__file__).resolve().parents[2] / "tools/push-challenges-template.xlsx"

    assert_table_filters_are_compatible(load_workbook(path))
