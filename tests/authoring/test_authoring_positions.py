import csv
import json

import pytest
from openpyxl import load_workbook

from authoring.csv_input import CSV_HEADERS, load_csv
from authoring.normalize import normalize
from authoring.preview import render_preview
from authoring.workbook import create_workbook
from authoring.xlsx_input import load_xlsx
from push_challenges_core.manifest import dump_manifest


@pytest.mark.parametrize("position", [None, 0, 20])
def test_workbook_position_reaches_export_and_preview(tmp_path, position):
    path = tmp_path / "positions.xlsx"
    create_workbook(path, author_rows=2)
    book = load_workbook(path)
    sheet = book["Challenges"]
    assert sheet["Q1"].value == "Position"
    sheet["A2"] = "A 1/4"
    sheet["C2"] = 100
    sheet["D2"] = "Web"
    sheet["Q2"] = position
    book.save(path)
    result = normalize(load_xlsx(path))
    assert result.manifest is not None, result.diagnostics
    assert result.manifest.challenges[0].position == position
    assert (
        f"Position: {position if position is not None else '(preserve/default)'}"
        in render_preview(result)
    )
    manifest_path = tmp_path / "positions.json"
    dump_manifest(result.manifest, manifest_path)
    assert (
        json.loads(manifest_path.read_text())["challenges"][0]["position"] == position
    )


@pytest.mark.parametrize("columns", [12, 16])
def test_older_workbook_layouts_preserve_position(tmp_path, columns):
    path = tmp_path / "legacy.xlsx"
    create_workbook(path, author_rows=2)
    book = load_workbook(path)
    sheet = book["Challenges"]
    sheet.delete_cols(columns + 1, sheet.max_column - columns)
    sheet["A2"] = "Legacy"
    sheet["C2"] = 100
    sheet["D2"] = "Web"
    book.save(path)
    result = normalize(load_xlsx(path))
    assert result.manifest is not None, result.diagnostics
    assert result.manifest.challenges[0].position is None


@pytest.mark.parametrize("position", ["-1", "1.5", "true", "2147483648"])
def test_invalid_position_blocks_authoring(tmp_path, position):
    path = tmp_path / "invalid.csv"
    headers = tuple(CSV_HEADERS[:21]) + ("Position",)
    values = dict.fromkeys(headers, "")
    values.update(
        Name="Bad", Value="100", Category="Web", Type="Standard", Position=position
    )
    values["Max Attempts"] = "0"
    with path.open("w", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerow(values)
    result = normalize(load_csv(path))
    assert result.manifest is None
    assert any(item.blocking for item in result.diagnostics)


@pytest.mark.parametrize("columns, position", [(16, None), (21, None), (22, 30)])
def test_csv_old_and_position_layouts(tmp_path, columns, position):
    path = tmp_path / "positions.csv"
    headers = (tuple(CSV_HEADERS[:21]) + ("Position",))[:columns]
    values = dict.fromkeys(headers, "")
    values.update(Name="A 3/4", Value="100", Category="Web", Type="Standard")
    values["Max Attempts"] = "0"
    if position is not None:
        values["Position"] = str(position)
    with path.open("w", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerow(values)
    result = normalize(load_csv(path))
    assert result.manifest is not None, result.diagnostics
    assert result.manifest.challenges[0].position == position
