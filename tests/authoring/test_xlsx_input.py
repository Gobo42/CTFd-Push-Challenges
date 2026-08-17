from openpyxl import load_workbook

from authoring.normalize import normalize
from authoring.workbook import create_workbook
from authoring.xlsx_input import load_xlsx


def test_xlsx_round_trip_joins_related_sheets(tmp_path):
    path = tmp_path / "authored.xlsx"
    create_workbook(path, author_rows=10)
    workbook = load_workbook(path)
    challenge = workbook["Challenges"]
    challenge.append(
        [
            "Example",
            "Prompt",
            100,
            "Web",
            "Visible",
            "Standard",
            "#beginner #web authentication",
            "Part 2",
            "Require All Flags",
            3,
            "Author",
            "nc host 31337",
        ]
    )
    flags = workbook["Flags"]
    flags.append(["Example", "Static", "flag{x}", "Case Sensitive"])
    flags.append(["Example", "Regex", r"flag\{.*\}", "Case Insensitive"])
    workbook["Solutions"].append(["Example", "Explain it", "Solved"])
    hints = workbook["Hints"]
    hints.append(["Example", "First", "Read it", 0, ""])
    hints.append(["Example", "Second", "Decode it", 5, "1"])
    workbook.save(path)

    result = normalize(load_xlsx(path))

    assert result.manifest is not None
    authored = result.manifest.challenges[0]
    assert authored.logic == "all"
    assert authored.max_attempts == 3
    assert tuple(flag.type for flag in authored.flags) == ("static", "regex")
    assert authored.solution.action == "upsert"
    assert authored.solution.state == "solved"
    assert authored.hints[1].required_hints == ("First",)
    assert authored.next == "Part 2"


def test_xlsx_orphan_hint_is_omitted_with_warning(tmp_path):
    path = tmp_path / "orphan-hint.xlsx"
    create_workbook(path, author_rows=5)
    workbook = load_workbook(path)
    workbook["Challenges"].append(
        ["Kept", "", 100, "Web", "Hidden", "Standard", "", "", "", 0, "", ""]
    )
    workbook["Hints"].append(["Deleted", "Old clue", "Ignore me", 0, ""])
    workbook.save(path)

    result = normalize(load_xlsx(path))

    assert result.manifest is not None
    assert result.manifest.challenges[0].hints == ()
    assert result.manifest.warnings[0].code == "orphan_hint_ignored"
    assert result.manifest.warnings[0].blocking is False


def test_xlsx_orphan_flag_and_duplicate_solution_are_blocking(tmp_path):
    path = tmp_path / "bad-related-rows.xlsx"
    create_workbook(path, author_rows=5)
    workbook = load_workbook(path)
    workbook["Challenges"].append(
        ["Kept", "", 100, "Web", "Hidden", "Standard", "", "", "", 0, "", ""]
    )
    workbook["Flags"].append(["Missing", "Static", "flag{x}", "Case Sensitive"])
    workbook["Solutions"].append(["Kept", "One", "Hidden"])
    workbook["Solutions"].append(["Kept", "Two", "Visible"])
    workbook.save(path)

    result = normalize(load_xlsx(path))

    assert result.manifest is None
    codes = {item.code for item in result.diagnostics}
    assert "orphan_flag" in codes
    assert "duplicate_solution_row" in codes
