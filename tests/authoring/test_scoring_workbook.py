import csv
import json

import pytest
from openpyxl import load_workbook

from authoring.csv_input import CSV_HEADERS, load_csv
from authoring.normalize import normalize
from authoring.preview import render_preview
from authoring.workbook import create_workbook
from authoring.xlsx_input import load_xlsx


def test_workbook_dynamic_bonus_rows_normalize_and_preview(tmp_path):
    path = tmp_path / "scoring.xlsx"
    create_workbook(path, author_rows=3)
    book = load_workbook(path)
    sheet = book["Challenges"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    values = {
        "Name": "Dynamic test",
        "Value": 100,
        "Category": "Web",
        "Type": "Dynamic",
        "Scoring Function": "Linear",
        "Minimum Value": 10,
        "Decay": 5,
        "Bonus Mode": "Percentage",
    }
    for key, value in values.items():
        sheet.cell(2, headers[key], value)
    for column, value in enumerate(("Dynamic test", 1, 1, 25.999), 1):
        book["Bonus Tiers"].cell(2, column, value)
    book.save(path)
    result = normalize(load_xlsx(path))
    assert result.manifest is not None, result.diagnostics
    spec = result.manifest.challenges[0]
    assert spec.scoring.function == "linear"
    assert spec.bonus.tiers[0].amount == "25.99"
    assert result.manifest.version == 2
    preview = render_preview(result)
    assert "Dynamic scoring: linear; initial 100; minimum 10; decay 5" in preview
    assert "Bonus Mode: percentage" in preview
    assert "Placements 1-1: 25.99" in preview


def test_old_workbook_without_new_columns_or_sheet_remains_supported(tmp_path):
    path = tmp_path / "legacy.xlsx"
    create_workbook(path, author_rows=2)
    book = load_workbook(path)
    sheet = book["Challenges"]
    sheet.delete_cols(13, sheet.max_column - 12)
    del book["Bonus Tiers"]
    sheet["A2"] = "Legacy"
    sheet["C2"] = 100
    sheet["D2"] = "Web"
    book.save(path)
    result = normalize(load_xlsx(path))
    assert result.manifest is not None, result.diagnostics
    assert result.manifest.challenges[0].bonus.mode == "preserve"


def test_new_csv_supports_dynamic_and_bonus_tiers(tmp_path):
    path = tmp_path / "scoring.csv"
    values = dict.fromkeys(CSV_HEADERS, "")
    values.update(
        {
            "Name": "Dynamic",
            "Value": "100",
            "Category": "Web",
            "Type": "Dynamic",
            "Max Attempts": "0",
            "Scoring Function": "Logarithmic",
            "Minimum Value": "20",
            "Decay": "10",
            "Bonus Mode": "Absolute",
            "Bonus Tiers": json.dumps([{"start": 1, "end": 2, "amount": "15"}]),
        }
    )
    with path.open("w", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=CSV_HEADERS)
        writer.writeheader()
        writer.writerow(values)
    result = normalize(load_csv(path))
    assert result.manifest is not None, result.diagnostics
    assert result.manifest.challenges[0].scoring.function == "logarithmic"
    assert result.manifest.challenges[0].bonus.tiers[0].end == 2


@pytest.mark.parametrize(
    ("mode", "challenge", "start", "end", "amount"),
    [
        ("Absolute", "Missing", 1, 1, 10),
        ("", "Example", 1, 1, 10),
        ("Disabled", "Example", 1, 1, 10),
        ("Absolute", "Example", 2, 3, 10),
        ("Absolute", "Example", 1, 1, 1.5),
        ("Percentage", "Example", 1, 1, -1),
    ],
)
def test_invalid_bonus_rows_block_workbook_export(
    tmp_path, mode, challenge, start, end, amount
):
    path = tmp_path / "invalid.xlsx"
    create_workbook(path, author_rows=2)
    book = load_workbook(path)
    book["Challenges"]["A2"] = "Example"
    book["Challenges"]["C2"] = 100
    book["Challenges"]["D2"] = "Web"
    book["Challenges"]["P2"] = mode
    for index, value in enumerate((challenge, start, end, amount), 1):
        book["Bonus Tiers"].cell(2, index, value)
    book.save(path)
    result = normalize(load_xlsx(path))
    assert result.manifest is None
    assert any(item.blocking for item in result.diagnostics)
